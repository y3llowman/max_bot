import asyncio
import io
import os
import sys
import unittest
from dataclasses import asdict, replace
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import AsyncMock, patch

TEST_DB = os.environ.get("TEST_DATABASE_URL")
if TEST_DB:
    os.environ["DATABASE_URL"] = TEST_DB
os.environ.setdefault("MAX_TOKEN", "test")
os.environ.setdefault("DATABASE_URL", "postgresql+asyncpg://max@localhost:5544/maxtest")
os.environ.setdefault("SECRET_KEY", "test-secret-key-" * 4)
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

from docx import Document
from httpx import ASGITransport, AsyncClient
from sqlalchemy import select
from sqlalchemy import text as sql

from app.api.depends import get_current_user
from data_fetching.msp_open_data import MspRecord
from databases import SessionLocal, engine, init_db, msp_registry
from databases.businesses_db import Business, save_business
from databases.users_db import User
from main import app
from notifications import worker
from notifications.models import BusinessProfile, Notification, RadarEvent
from notifications.planner import MSK

TODAY = date(2026, 9, 26)
MORNING = datetime(2026, 9, 26, 9, 0, tzinfo=MSK)
INN = "7707083893"
MAX_USER = 111
BOT_APP = {"web_app": "radar_bot", "contact_id": 999}
RECORD = MspRecord(
    name='ООО "СЕВЕРНЫЙ ВЕТЕР"', subject_type="UL", category=1, ogrn="1027700132195", inn=INN,
    main_activity_code="41.20", main_activity_name="Строительство жилых и нежилых зданий", region_code="16",
    is_new=False, date_registered="10.08.2016 00:00:00", date_excluded=None, phone=None, email=None, website=None, employees_num=None,
    has_licenses=False, is_hitech=False, is_partnership=False, is_social=False,
)


@unittest.skipUnless(TEST_DB, "нужен TEST_DATABASE_URL")
class RadarDbTest(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        async with engine.begin() as conn:
            await conn.execute(sql("DROP SCHEMA public CASCADE"))
            await conn.execute(sql("CREATE SCHEMA public"))
        await init_db()
        async with SessionLocal() as db:
            user = User(max_user_id=MAX_USER, notification_settings={"quiet": False})
            db.add(user)
            await db.flush()
            await save_business(db, user.id, RECORD)
        await msp_registry.load(iter([("10.09.2026", RECORD)]), "full", "test")
        self.send = AsyncMock(return_value="mid-1")
        self.deleted = AsyncMock()
        patches = [patch.object(worker, "send_html", self.send),
                   patch.object(worker, "bot_app", AsyncMock(return_value=BOT_APP)),
                   patch.object(worker.bot, "delete_message", self.deleted)]
        for p in patches:
            p.start()
            self.addCleanup(p.stop)

    async def asyncTearDown(self):
        await engine.dispose()

    async def api(self, max_user_id: int = MAX_USER) -> AsyncClient:
        async with SessionLocal() as db:
            user = (await db.execute(select(User).where(User.max_user_id == max_user_id))).scalar_one()
        app.dependency_overrides[get_current_user] = lambda: user
        self.addCleanup(app.dependency_overrides.clear)
        client = AsyncClient(transport=ASGITransport(app=app), base_url="http://test")
        self.addAsyncCleanup(client.aclose)
        return client

    async def reply(self):
        await worker.answered(MAX_USER, "mid-1")

    async def answer(self, regime: str | None, headcount: int | None):
        async with SessionLocal() as db:
            profile = await db.get(BusinessProfile, INN) or BusinessProfile(inn=INN, flags={}, bank_biks=[])
            profile.tax_regime, profile.headcount = regime, headcount
            profile.has_employees = bool(headcount)
            db.add(profile)
            await db.commit()
        return await worker.materialize_for(INN, TODAY)

    async def events(self, **where) -> list[RadarEvent]:
        async with SessionLocal() as db:
            query = select(RadarEvent).filter_by(inn=INN, **where).order_by(RadarEvent.due, RadarEvent.id)
            return list(await db.scalars(query))

    async def notifications(self) -> list[Notification]:
        async with SessionLocal() as db:
            return list(await db.scalars(select(Notification).order_by(Notification.id)))

    async def test_profile_answers_change_obligations(self):
        added, removed = await worker.materialize_for(INN, TODAY)
        self.assertEqual((added, removed), (["Бухгалтерская отчётность"], []))

        added, removed = await self.answer("usn_income", 36)
        self.assertIn("Сведения о выполнении квоты для инвалидов", added)
        self.assertIn("Декларация по УСН", added)
        total = len(await self.events())
        await worker.materialize_for(INN, TODAY)
        self.assertEqual(len(await self.events()), total, "повторный запуск не плодит события")

        rsv = [e for e in await self.events() if e.payload["code"] == "rsv"][0]
        self.assertEqual((rsv.due, rsv.payload["shifted"], rsv.payload["original"]), (date(2026, 10, 26), True, "2026-10-25"))

        added, removed = await self.answer("usn_income", 0)
        self.assertEqual(added, [])
        self.assertIn("Расчёт по страховым взносам (РСВ)", removed)
        self.assertFalse([e for e in await self.events() if e.payload.get("code") == "rsv"])

    async def test_reminders_are_grouped_and_buttons_close_tasks(self):
        await self.answer("usn_income", 36)
        await worker.queue_reminders(MORNING)
        await worker.queue_reminders(MORNING)
        pending = await self.notifications()
        self.assertTrue(pending)
        self.assertTrue(all(n.label == "T-30" for n in pending))

        await worker.dispatch()
        self.send.assert_awaited_once()
        _, text, kb = self.send.await_args.args
        self.assertIn(" дней: 26 октября 2026", text)
        self.assertIn("• Расчёт по страховым взносам (РСВ) — за 9 месяцев 2026", text)
        self.assertNotIn("почему вам", text, "в чат — коротко, подробности — в приложении")
        self.assertTrue(all(n.status == "sent" for n in await self.notifications()))

        ids = [n.event_id for n in pending]
        joined = ",".join(map(str, ids))
        self.assertEqual(kb["payload"]["buttons"], [
            [{"type": "callback", "text": "📋 В список дел", "payload": f"ev:{joined}:list"},
             {"type": "callback", "text": "✅ Сделано", "payload": f"ev:{joined}:done"}],
            [{"type": "callback", "text": "🔕 Не актуально", "payload": f"ev:{joined}:mute"},
             {"type": "open_app", "text": "Открыть", **BOT_APP, "payload": f"task_{ids[0]}"}],
        ])
        self.assertEqual(await worker.apply_action(MAX_USER, ids, "list"), "Добавили в список дел — он в мини-приложении")
        self.assertTrue(all(e.in_list and e.status == "open" for e in await self.events() if e.id in ids),
                        "в списке дел задача открыта — напоминания идут, пока её не закроют")

        self.assertEqual(await worker.apply_action(MAX_USER, ids, "done"), "Отмечено как выполненное")
        self.assertEqual(await worker.apply_action(222, ids, "done"), "Задача не найдена", "чужие задачи не трогаем")
        async with SessionLocal() as db:
            user = (await db.execute(select(User))).scalar_one()
            for event in await self.events():
                if event.id in ids:
                    await worker._queue(db, user.id, event, "T-1", MORNING)
            await db.commit()
        self.send.reset_mock()
        await self.reply()
        self.send.assert_not_awaited()
        self.assertTrue(all(n.status == "cancelled" for n in await self.notifications() if n.label == "T-1"))

    async def test_queue_sends_one_message_until_answered(self):
        await self.answer("usn_income", 36)
        deadlines = await self.events(type="deadline")
        a = deadlines[0]
        b = next(e for e in deadlines if e.due != a.due)
        async with SessionLocal() as db:
            user = (await db.execute(select(User))).scalar_one()
            for label in ("T-7", "T-1"):
                await worker._queue(db, user.id, a, label, MORNING)
            await worker._queue(db, user.id, b, "T-1", MORNING)
            await db.commit()
        await worker.dispatch()
        self.send.assert_awaited_once()
        self.assertIn("Ещё 1 сообщение", self.send.await_args.args[1])
        self.assertEqual({n.label: n.status for n in await self.notifications() if n.event_id == a.id},
                         {"T-7": "cancelled", "T-1": "sent"})
        await worker.dispatch()
        self.send.assert_awaited_once()
        async with SessionLocal() as db:
            user = (await db.execute(select(User))).scalar_one()
            user.awaiting_since -= timedelta(hours=25)
            await db.commit()
        await worker.dispatch()
        self.assertEqual(self.send.await_count, 2)

    async def test_overdue_is_reminded_every_day_until_closed(self):
        await self.answer(None, None)
        buh = (await self.events())[0]
        async with SessionLocal() as db:
            (await db.get(RadarEvent, buh.id)).due = date(2026, 9, 20)
            await db.commit()
        await worker.queue_reminders(MORNING)
        await worker.queue_reminders(datetime(2026, 9, 27, 9, 0, tzinfo=MSK))
        self.assertEqual([n.label for n in await self.notifications()], ["overdue:260926", "overdue:260927"])

    async def test_chat_disabled_cancels_reminders(self):
        await self.answer("usn_income", 36)
        async with SessionLocal() as db:
            user = (await db.execute(select(User))).scalar_one()
            user.notification_settings = {"chat": False}
            await db.commit()
        await worker.queue_reminders(MORNING)
        await worker.dispatch()
        self.send.assert_not_awaited()

    async def test_msp_exclusion_and_category_change(self):
        await worker.registry_loaded(INN, asdict(RECORD))
        self.assertEqual(await self.events(type="msp.excluded"), [])

        excluded = asdict(replace(RECORD, date_excluded="10.07.2027 00:00:00"))
        self.assertEqual(len(await worker.process_msp(INN, excluded, TODAY)), 1)
        self.assertEqual(await worker.process_msp(INN, excluded, TODAY), [], "открытое состояние не дублируется")
        await worker.dispatch()
        self.assertIn("Компания исключена из реестра МСП", self.send.await_args.args[1])
        self.assertIn("10 июля 2027", self.send.await_args.args[1])

        await worker.process_msp(INN, asdict(RECORD), TODAY)
        self.assertIsNotNone((await self.events(type="msp.excluded"))[0].resolved_at)

        grown = asdict(replace(RECORD, category=2))
        [event] = await worker.process_msp(INN, grown, TODAY)
        self.assertEqual((event.type, event.payload["lost_micro"], event.due), ("msp.category_changed", True, date(2027, 1, 26)))

    async def test_open_condition_repeats_and_done_is_rechecked(self):
        excluded = asdict(replace(RECORD, date_excluded="10.07.2027 00:00:00"))
        [event] = await worker.process_msp(INN, excluded, TODAY)
        await worker.dispatch()
        async with SessionLocal() as db:
            (await db.get(RadarEvent, event.id)).last_notified_at = MORNING - timedelta(days=30)
            await db.commit()
        await worker.queue_reminders(MORNING)
        self.assertIn("repeat:260926", [n.label for n in await self.notifications()])
        await self.reply()

        self.assertEqual(await worker.apply_action(MAX_USER, [event.id], "done"),
                         "Проверим по реестру: если запись останется, напомним")
        self.assertEqual(await worker.process_msp(INN, excluded, TODAY), [], "реестр обновляется не сразу — ждём")
        async with SessionLocal() as db:
            (await db.get(RadarEvent, event.id)).last_notified_at = MORNING - timedelta(days=30)
            await db.commit()
        [reopened] = await worker.process_msp(INN, excluded, TODAY)
        self.assertEqual((reopened.id, reopened.status), (event.id, "open"))
        self.assertIn("recheck:260926", [n.label for n in await self.notifications()])

    async def test_new_company_replaces_the_old_one(self):
        await self.answer("usn_income", 5)
        await worker.demo_remind(MAX_USER)
        async with SessionLocal() as db:
            user = (await db.execute(select(User))).scalar_one()
            event = (await self.events())[-1]
            await worker._queue(db, user.id, event, "T-1", MORNING + timedelta(days=365))
            await db.commit()
            other = replace(RECORD, inn="1650273744", name='ООО "КАМСКАЯ"')
            await save_business(db, user.id, other)
            self.assertEqual((await worker.current_business(db, user.id)).inn, other.inn)
            self.assertEqual(await worker.user_events(db, MAX_USER, [event.id]), (user, []), "старая компания отвязана")
        self.assertEqual([n.status for n in await self.notifications() if n.label == "T-1"], ["cancelled"])

    async def test_background_scan_runs_once(self):
        started = []

        async def slow(inns, today):
            started.append(inns)
            await asyncio.sleep(0.05)

        knm = AsyncMock(side_effect=slow)
        worker._erknm_checked.clear()
        with patch.object(worker, "process_inspections", knm):
            await asyncio.gather(worker.scan_in_background(INN), worker.scan_in_background(INN))
            self.assertEqual(len(started), 1, "пока идёт проверка, вторая не запускается")
            await worker.scan_in_background(INN)
        self.assertEqual(knm.await_count, 1, "ЕРКНМ по одному ИНН — раз в день")

    async def test_inspections_from_erknm(self):
        from data_fetching.erknm_client import Knm
        check = Knm(erpid="1001", inn=INN, classification="КНМ", status="Ожидает проведения",
                    status_key="TYPE_WAITING_CARRY_OUT", kind="Документарная проверка", control="Пожарный надзор",
                    authority="ГУ МЧС", type_name="Плановое КНМ", start="2026-10-19", stop="2026-10-21", warning=None)
        month = lambda found: (lambda year, m, inns: found if (year, m) == (2026, 10) else [])
        with patch.object(worker.erknm_client, "fetch_month", month([check])):
            await worker.process_inspections({INN}, TODAY)
            await worker.process_inspections({INN}, TODAY)
        [event] = await self.events(type="inspection.planned")
        self.assertEqual((event.due, event.source), (date(2026, 10, 19), "erknm"))
        await worker.dispatch()
        self.send.assert_awaited_once()
        self.assertIn("ГУ МЧС: документарную проверку с 19 октября 2026", self.send.await_args.args[1])

        cancelled = replace(check, status="Отменено", status_key="TYPE_CANCELED")
        with patch.object(worker.erknm_client, "fetch_month", month([cancelled])):
            await worker.process_inspections({INN}, TODAY)
        self.assertEqual((await self.events(type="inspection.planned"))[0].status, "muted")

    async def test_demo_inspection(self):
        self.assertIn("Готово", await worker.demo_event(MAX_USER, "knm"))
        text = self.send.await_args.args[1]
        self.assertIn("Запланирована проверка", text)
        self.assertNotIn("Ваши права", text, "подробности — на экране события")
        [event] = await self.events(type="inspection.planned")
        client = await self.api()
        details = (await client.get(f"/api/tasks/{event.id}")).json()
        sections = {section["title"]: section for section in details["sections"]}
        self.assertEqual(sections["Ваши права"]["id"], "law")
        self.assertIn("Что сделать", sections)
        self.assertIn("Имитация для демо", "\n".join(line for s in details["sections"] for line in s.get("body", [])))

    async def test_new_laws_reach_matching_company(self):
        await self.answer("usn_income", 5)
        day = TODAY - timedelta(days=1)

        def act(eo, header, name, block):
            return {"eoNumber": eo, "complexName": f'{header}\n "{name}"', "name": f'"{name}"', "block": block,
                    "publishDateShort": f"{day}T00:00:00", "pagesCount": 2, "pdfFileLength": 50_000}

        acts = [
            act("0001202609250001", "Приказ Федеральной налоговой службы от 20.08.2026 № ЕД-7-3/555@",
                "Об утверждении формы налоговой декларации по налогу, уплачиваемому в связи с применением "
                "упрощенной системы налогообложения, и порядка ее заполнения", "federal_authorities"),
            act("0001202609250002", "Федеральный закон от 24.09.2026 № 400-ФЗ",
                "О внесении изменений в Федеральный закон «О применении контрольно-кассовой техники»", "president"),
            act("0001202609250003", "Приказ Федеральной антимонопольной службы от 01.09.2026 № 70/26",
                "Об утверждении Порядка представления сведений о доходах, расходах, об имуществе и обязательствах "
                "имущественного характера в ФАС России и ее территориальных органах", "federal_authorities"),
            act("0001202609250004", "Федеральный закон от 24.09.2026 № 401-ФЗ",
                "О внесении изменений в отдельные законодательные акты Российской Федерации", "president"),
            act("0001202609250005", "Закон Амурской области от 24.09.2026 № 820-ОЗ",
                "Об установлении налоговой ставки по налогу, взимаемому в связи с применением упрощенной системы "
                "налогообложения", "subjects"),
        ]
        pdf = b"%PDF-1.4 test"
        with patch.object(worker.pravo_client, "published", lambda d: acts if d == day else []), \
                patch.object(worker.pravo_client, "find_act", lambda *args: None), \
                patch.object(worker.pravo_client, "pdf", lambda eo: pdf), \
                patch.object(worker.bot, "send_message", AsyncMock()) as send_file:
            worker._law_pdf.cache_clear()
            self.assertEqual(await worker.ingest_laws(TODAY), 3)
            [law] = await self.events(type="law.upcoming")
            [question] = await self.events(type="profile.question")
            self.assertEqual((law.key, law.due), ("law:0001202609250001:" + INN, None))
            self.assertEqual(question.payload["flag"], "cash_register")

            await worker.dispatch()
            self.send.assert_awaited_once()
            self.assertIn("Ещё 1 сообщение — пришлю по одному", self.send.await_args.args[1])
            await self.reply()
            texts_sent = {call.args[1].split("\n")[0]: call.args for call in self.send.await_args_list}
            self.assertEqual(len(texts_sent), 2)
            _, text, kb = next(args for title, args in texts_sent.items() if "Вышел акт" in title)
            self.assertIn("Почему вам: вы на УСН доходы.", text)
            buttons = [b for row in kb["payload"]["buttons"] for b in row]
            self.assertEqual([b.get("payload") for b in buttons],
                             [f"ev:{law.id}:list", f"ev:{law.id}:done", f"ev:{law.id}:mute", f"task_{law.id}", None])
            self.assertEqual((buttons[-1]["type"], buttons[-1]["url"]),
                             ("link", "http://publication.pravo.gov.ru/document/0001202609250001"))
            _, text, kb = next(args for title, args in texts_sent.items() if "кассу" in title)
            self.assertEqual([b["payload"] for b in kb["payload"]["buttons"][0]],
                             ["flag:cash_register:yes", "flag:cash_register:no"])
            self.assertEqual(send_file.await_count, 0, "PDF не скачиваем, пока не включён LAW_PDF")

            with patch.object(worker, "LAW_PDF", True):
                self.assertEqual(await worker.answer_flag(MAX_USER, "cash_register", True), "Запомнили — присылаем, что вышло")
                await self.reply()
            [kkt] = [e for e in await self.events(type="law.upcoming") if e.payload["eo"] == "0001202609250002"]
            self.assertIsNone(kkt.due, "дата вступления в силу — только из текста, а текста в документированном API нет")
            self.assertEqual((await self.events(type="profile.question"))[0].status, "done")
            [pdf_call] = send_file.await_args_list
            media = pdf_call.kwargs["attachments"][0]
            self.assertEqual((media.buffer, media.filename),
                             (pdf, "Федеральный закон от 24.09.2026 N 400-ФЗ.pdf"))

            self.assertEqual(await worker.ingest_laws(TODAY), 0, "разобранные акты второй раз не рассылаются")
            self.assertEqual(await worker.answer_flag(MAX_USER, "marked_goods", False),
                             "Запомнили: такие акты присылать не будем")
        async with SessionLocal() as db:
            self.assertEqual((await db.get(BusinessProfile, INN)).flags, {"cash_register": True, "marked_goods": False})

        async with SessionLocal() as db:
            user = (await db.execute(select(User))).scalar_one()
        app.dependency_overrides[get_current_user] = lambda: user
        self.addCleanup(app.dependency_overrides.clear)
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            details = (await client.get(f"/api/tasks/{kkt.id}")).json()
        self.assertEqual([s["id"] for s in details["sections"]], ["summary", "why", "steps", "law"])
        self.assertEqual(details["sections"][1]["body"], ["Вы принимаете оплату через кассу"])

    async def test_connect_in_miniapp_continues_in_chat(self):
        from app.api.routes import company as company_routes
        from bot import message_handler as bot_flow
        async with SessionLocal() as db:
            user = (await db.execute(select(User))).scalar_one()
        app.dependency_overrides[get_current_user] = lambda: user
        self.addCleanup(app.dependency_overrides.clear)
        next_step, scan = AsyncMock(), AsyncMock()
        with patch.object(bot_flow, "next_step", next_step), patch.object(company_routes, "scan_in_background", scan):
            async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
                self.assertEqual((await client.post("/api/session", json={"inn": INN})).status_code, 200)
        next_step.assert_awaited_once_with(MAX_USER)
        scan.assert_awaited_once_with(INN)

        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            missing = await client.post("/api/session", json={"inn": "7743212897"})
        self.assertEqual((missing.status_code, missing.json()["detail"]), (404, "INN not found in the SME registry"))

    async def test_demo_tour_covers_every_step_and_cleans_up(self):
        from notifications import demo as tour
        await self.answer("usn_income", 40)
        async with SessionLocal() as db:
            db.add(worker.LawRecord(eo_number="0001202609250001", header="Федеральный закон от 24.09.2026 № 400-ФЗ",
                                    name="О внесении изменений в Федеральный закон «О применении контрольно-кассовой техники»",
                                    published=TODAY - timedelta(days=2), topics=["kkt", "usn"], amended=[], effective=[]))
            await db.commit()
        cards, shown = [], AsyncMock()
        card = AsyncMock(side_effect=lambda uid, text, kb=None: cards.append((text, kb)))
        with patch.object(tour, "send_html", card), patch.object(tour, "bot_app", AsyncMock(return_value=BOT_APP)), \
                patch.object(worker.bot, "send_message", AsyncMock()) as files, \
                patch.object(worker.pravo_client, "pdf", lambda eo: b"%PDF-1.4"):
            worker._law_pdf.cache_clear()
            await worker.hold(MAX_USER, await tour.intro(MAX_USER))
            self.assertIn(f"Демо всех функций: {len(tour.STEPS)} шагов", cards[0][0])
            arg = cards[0][1]["payload"]["buttons"][0][0]["payload"].removeprefix("demo:")
            waited = False
            for n in range(1, len(tour.STEPS) + 1):
                await worker.answered(MAX_USER, "mid-1", release=False)
                sent_before = self.send.await_count
                await tour.run(MAX_USER, arg, shown)
                for _ in range(10):
                    text, kb = self.send.await_args.args[1:]
                    if f"Шаг {n} из {len(tour.STEPS)}" in text:
                        break
                    waited |= "пришлю по одному" in text
                    await self.reply()
                self.assertIn(f"Шаг {n} из {len(tour.STEPS)}", text)
                self.assertNotIn("⚠️", text, f"шаг {n}: {text}")
                if tour.STEPS[n - 1].run not in (None, tour.profile):
                    self.assertGreater(self.send.await_count, sent_before + 1, f"шаг {n} ничего не прислал")
                self.assertEqual(tour.running(MAX_USER), n < len(tour.STEPS))
                arg = kb["payload"]["buttons"][-1][0]["payload"].removeprefix("demo:")
            self.assertTrue(waited, "у шагов с несколькими сообщениями есть «Ещё N — пришлю по одному»")
            shown.assert_awaited_once()
            self.assertEqual(kb["payload"]["buttons"][-1][0]["payload"], "demo:clean")
            self.assertEqual(files.await_count, 0, "PDF закона не качаем без LAW_PDF; документы — по кнопке, не пачкой")

            texts = "\n".join(call.args[1] for call in self.send.await_args_list)
            for expected in ("Запланирована проверка", "Запланирован профилактический визит", "предостережение",
                             "Недостоверные сведения в ЕГРЮЛ", "Отметку о недостоверности сняли", "дисквалифицирован",
                             "Дисквалификация руководителя закончилась", "Компания в процессе прекращения",
                             "В ЕГРЮЛ изменились сведения: руководитель", "Уведомите МВД", "исключена из реестра МСП",
                             "Изменилась категория МСП", "Компании нет в реестре МСП", "Вышел акт, который вас касается",
                             "Вы принимаете оплату через кассу"):
                self.assertIn(expected, texts)

            demo_events = await self.events(source="demo")
            self.assertTrue(demo_events)
            await tour.run(MAX_USER, "clean", shown)
        self.assertIn(f"Убрали демо-события: {len(demo_events)}", cards[-1][0])
        self.assertEqual(await self.events(source="demo"), [])
        self.assertTrue(await self.events(type="deadline"), "настоящие сроки остались")

    async def test_only_one_demo_tour_at_a_time(self):
        from notifications import demo as tour
        await self.answer("usn_income", 40)
        cards = []
        card = AsyncMock(side_effect=lambda uid, text, kb=None: cards.append((text, kb)))
        with patch.object(tour, "send_html", card), patch.object(tour, "bot_app", AsyncMock(return_value=BOT_APP)):
            await tour.intro(MAX_USER)
            first = cards[-1][1]["payload"]["buttons"][0][0]["payload"].removeprefix("demo:")
            await tour.run(MAX_USER, first, AsyncMock())
            await tour.run(MAX_USER, first.replace(":1", ":2"), AsyncMock())
            self.assertTrue(tour.running(MAX_USER))
            self.assertTrue(any(n.status == "pending" for n in await self.notifications()))

            await tour.intro(MAX_USER)
            self.assertFalse(tour.running(MAX_USER))
            self.assertFalse([n for n in await self.notifications() if n.status == "pending"], "очередь прошлого тура отменена")
            second = cards[-1][1]["payload"]["buttons"][0][0]["payload"].removeprefix("demo:")
            self.assertNotEqual(first, second)

            sent = self.send.await_count
            await tour.run(MAX_USER, first.replace(":1", ":3"), AsyncMock())
            self.assertIn("кнопка прошлого демо", cards[-1][0])
            self.assertEqual(self.send.await_count, sent, "шаг прошлого тура не выполнен")
            self.assertFalse(tour.running(MAX_USER))

            await tour.run(MAX_USER, second, AsyncMock())
            self.assertTrue(tour.running(MAX_USER))
            await tour.run(MAX_USER, "stop", AsyncMock())
            self.assertFalse(tour.running(MAX_USER))

    async def test_miniapp_api_and_document(self):
        async with SessionLocal() as db:
            user = (await db.execute(select(User))).scalar_one()
        app.dependency_overrides[get_current_user] = lambda: user
        self.addCleanup(app.dependency_overrides.clear)
        client = AsyncClient(transport=ASGITransport(app=app), base_url="http://test")
        self.addAsyncCleanup(client.aclose)

        company = (await client.get("/api/session")).json()
        self.assertEqual((company["region"], company["category"], company["needsAnswers"]),
                         ("Республика Татарстан", "Микропредприятие", True))
        await self.answer("usn_income", 36)
        company = (await client.get("/api/company")).json()
        self.assertEqual((company["headcount"], company["needsAnswers"]), ("36–100 человек", False))
        self.assertEqual([b["title"][:22] for b in company["benefits"]], ["Декларации можно сдава", "Можно применять УСН: с"])
        self.assertNotIn("counterparties", company)

        self.assertFalse((await client.get("/api/dashboard")).json()["unread"])
        await worker.demo_remind(MAX_USER)
        self.assertTrue((await client.get("/api/dashboard")).json()["unread"])
        self.assertEqual((await client.post("/api/settings/notifications/seen")).status_code, 204)
        self.assertFalse((await client.get("/api/dashboard")).json()["unread"])

        dashboard = (await client.get("/api/dashboard")).json()
        quota = next(t for t in dashboard["tasks"] if t["title"].startswith("Сведения о выполнении квоты"))
        self.assertEqual(quota["periodicity"], "Ежемесячно")
        details = (await client.get(f"/api/tasks/{quota['id']}")).json()
        self.assertEqual([s["id"] for s in details["sections"]], ["why", "steps", "how", "law", "risks"])
        self.assertTrue(details["document"])
        self.assertEqual(details["heading"], "Сведения о выполнении квоты для инвалидов за сентябрь 2026")
        law = next(s for s in details["sections"] if s["id"] == "law")
        self.assertTrue(law["link"]["url"].startswith("https://www.consultant.ru/"))

        with patch.object(worker.bot, "send_message", AsyncMock()) as send_message:
            self.assertEqual((await client.post(f"/api/tasks/{quota['id']}/document")).status_code, 204)
        media = send_message.await_args.kwargs["attachments"][0]
        doc = Document(io.BytesIO(media.buffer))
        self.assertIn("Приказ", "\n".join(p.text for p in doc.paragraphs))
        self.assertIn("Республика Татарстан", "\n".join(p.text for p in doc.paragraphs))

        rsv = next(t for t in dashboard["tasks"] if t["title"].startswith("Расчёт по страховым"))
        with patch.object(worker.bot, "send_message", AsyncMock()) as send_message:
            self.assertEqual((await client.post(f"/api/tasks/{rsv['id']}/document")).status_code, 204)
        doc = Document(io.BytesIO(send_message.await_args.kwargs["attachments"][0].buffer))
        self.assertIn("1151111", [c.text for t in doc.tables for row in t.rows for c in row.cells])
        self.assertIn("Памятка к отчёту", send_message.await_args.kwargs["text"])

        settings = (await client.get("/api/settings/notifications")).json()
        self.assertEqual(settings["remind"], "d30-7-1")
        settings |= {"remind": "d3-0", "quietRange": "23:00–07:00"}
        self.assertEqual((await client.put("/api/settings/notifications", json=settings)).status_code, 204)
        self.assertEqual((await client.get("/api/settings/notifications")).json()["quietRange"], "23:00–07:00")
        self.assertEqual((await client.put("/api/settings/notifications", json={"remind": "d99"})).status_code, 422)

    async def test_notice_document_has_requisites(self):
        await self.answer("usn_ie", 1)
        notice = next(e for e in await self.events() if e.payload.get("code") == "usn_notice")
        async with SessionLocal() as db:
            answers = await db.get(BusinessProfile, INN)
            answers.kpp, answers.address = "165001001", "423800, г. Набережные Челны, пр. Мира, д. 1"
            await db.commit()
        async with SessionLocal() as db:
            with patch.object(worker.bot, "send_message", AsyncMock()) as send_message:
                await worker.send_document(db, notice, MAX_USER)
        doc = Document(io.BytesIO(send_message.await_args.kwargs["attachments"][0].buffer))
        cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
        self.assertIn(INN, cells)
        self.assertIn("165001001", cells)
        self.assertIn("423800, г. Набережные Челны, пр. Мира, д. 1", cells)
        self.assertIn("18210501021011000110", cells)

    async def test_registry_slice_is_honest_about_missing_companies(self):
        from bot import message_handler as bot_flow
        from data_fetching import msp_open_data
        await msp_registry.load(msp_open_data.read_slice(), "slice", msp_open_data.SLICE.name)
        self.assertFalse(await msp_registry.is_complete())
        manilla = await msp_registry.find("7743212897")
        self.assertEqual((manilla.subject_type, manilla.main_activity_code[:5], manilla.region_code), ("UL", "56.10", "77"))
        self.assertIsNotNone(await msp_registry.find("3666155612"))
        self.assertIsNone(await msp_registry.find("7707083893"))
        text = await bot_flow.not_in_registry("7707083893")
        self.assertIn("нет в срезе реестра МСП", text)
        self.assertIn("срез открытых данных ФНС на 10.09.2026", text)

        client = await self.api()
        missing = await client.post("/api/session", json={"inn": "7707083893"})
        self.assertEqual((missing.status_code, missing.json()["detail"]),
                         (404, "INN not found in the SME registry slice loaded in this MVP"))
        from app.api.routes import company as company_routes
        with patch.object(company_routes, "scan_in_background", AsyncMock()):
            self.assertEqual((await client.post("/api/company/refresh")).status_code, 200)
        self.assertEqual(await self.events(source="msp"), [], "компании нет в срезе — это не «исключена из реестра»")
        company = (await client.get("/api/company")).json()
        self.assertIn("срез открытых данных ФНС", company["source"]["name"])

    async def test_migrations_match_models_and_adopt_old_database(self):
        from alembic.autogenerate import compare_metadata
        from alembic.migration import MigrationContext
        from databases import Base

        def diff(connection):
            return compare_metadata(MigrationContext.configure(connection, opts={"compare_type": True}), Base.metadata)

        async with engine.connect() as conn:
            self.assertEqual(await conn.run_sync(diff), [])
        async with engine.begin() as conn:
            for statement in ("DROP TABLE alembic_version", "DROP TABLE msp_registry", "DROP TABLE msp_registry_loads",
                              "ALTER TABLE business_profiles DROP COLUMN kpp, DROP COLUMN address, "
                              "DROP COLUMN director_position, DROP COLUMN director_name",
                              "ALTER TABLE users ALTER COLUMN max_user_id TYPE integer"):
                await conn.execute(sql(statement))
        await init_db()
        await init_db()
        async with engine.connect() as conn:
            self.assertEqual(await conn.run_sync(diff), [])
            self.assertEqual((await conn.execute(sql("SELECT version_num FROM alembic_version"))).scalar(), "0003")
        async with SessionLocal() as db:
            db.add(User(max_user_id=9_000_000_000_000_000_001))
            await db.commit()
        self.assertEqual(len(await self.events()), 0)

    async def test_api_auth_and_input_errors(self):
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            self.assertEqual((await client.get("/api/dashboard")).status_code, 401)
            self.assertEqual((await client.get("/api/dashboard", headers={"Authorization": "Bearer nope"})).status_code, 401)
            self.assertEqual((await client.post("/api/user/auth", json={"initData": "auth_date=1&hash=x"})).status_code, 401)
            from sign_init_data import TEST_USER, sign
            auth = await client.post("/api/user/auth", json={"initData": sign(TEST_USER)})
            headers = {"Authorization": f"Bearer {auth.json()['access_token']}"}
            self.assertEqual((await client.get("/api/user/me", headers=headers)).json()["id"], TEST_USER)
            self.assertEqual((await client.get("/api/session", headers=headers)).status_code, 404)
            forged = sign(TEST_USER).replace("api_check", "admin")
            self.assertEqual((await client.post("/api/user/auth", json={"initData": forged})).status_code, 401)
        client = await self.api()
        self.assertEqual((await client.post("/api/session", json={"inn": "7707083890"})).status_code, 422)
        self.assertEqual((await client.post("/api/session", json={"inn": "123"})).status_code, 422)

        from app.api.routes import company as company_routes
        async with SessionLocal() as db:
            await db.execute(sql("DELETE FROM msp_registry WHERE inn = :inn"), {"inn": INN})
            await db.commit()
        with patch.object(company_routes, "scan_in_background", AsyncMock()) as scan:
            company = await client.post("/api/company/refresh")
        scan.assert_awaited_once_with(INN)
        self.assertEqual((company.status_code, company.json()["inn"]), (200, INN))
        self.assertEqual([e.type for e in await self.events(source="msp")], ["msp.not_found"])

    async def test_demo_deadlines_open_in_app_and_are_labelled(self):
        import seed_demo_tasks
        await seed_demo_tasks.main(INN)
        client = await self.api()
        tasks = (await client.get("/api/dashboard")).json()["tasks"]
        demo = [t for t in tasks if t["subtitle"].startswith("Демо")]
        self.assertEqual(len(demo), 5)
        for task in demo:
            details = await client.get(f"/api/tasks/{task['id']}")
            self.assertEqual(details.status_code, 200, details.text)
            self.assertIn("демо-данные", " ".join(" ".join(s.get("body") or []) for s in details.json()["sections"]))
        from notifications import demo as tour
        await tour.cleanup(INN)
        self.assertEqual(await self.events(source="demo"), [])

    async def test_list_feed_and_closing_in_app(self):
        [event] = await worker.process_msp(INN, None, TODAY)
        await worker.dispatch()
        client = await self.api()
        dashboard = (await client.get("/api/dashboard")).json()
        self.assertNotIn(str(event.id), [t["id"] for t in dashboard["tasks"]], "без срока и не в списке — не задача")
        [item] = (await client.get("/api/feed")).json()
        self.assertEqual((item["id"], item["title"], item["status"], item["listed"], item["severity"]),
                         (str(event.id), "Компании нет в реестре МСП", "open", False, "warning"))

        self.assertEqual((await client.post(f"/api/tasks/{event.id}/list")).status_code, 204)
        self.deleted.assert_awaited_once_with("mid-1")
        task = next(t for t in (await client.get("/api/dashboard")).json()["tasks"] if t["id"] == str(event.id))
        self.assertNotIn("due", task)
        self.assertTrue((await client.get(f"/api/tasks/{event.id}")).json()["listed"])
        self.assertTrue((await client.get("/api/feed")).json()[0]["listed"])
        self.assertEqual((await client.post("/api/tasks/999999/list")).status_code, 404)

        await self.answer("usn_income", 36)
        far = next(e for e in await self.events(type="deadline") if e.due > TODAY + timedelta(days=30))
        task = next(t for t in (await client.get("/api/dashboard")).json()["tasks"] if t["id"] == str(far.id))
        self.assertFalse(task["listed"])
        self.assertEqual((await client.post(f"/api/tasks/{far.id}/list")).status_code, 204)
        task = next(t for t in (await client.get("/api/dashboard")).json()["tasks"] if t["id"] == str(far.id))
        self.assertTrue(task["listed"])

    async def test_blocked_bot_stops_messages_until_return(self):
        from maxapi.exceptions.max import MaxApiError
        await self.answer("usn_income", 36)
        deadlines = await self.events(type="deadline")
        later = next(e for e in deadlines if e.due != deadlines[0].due)
        async with SessionLocal() as db:
            user = (await db.execute(select(User))).scalar_one()
            for event in (deadlines[0], later):
                await worker._queue(db, user.id, event, "T-1", MORNING)
            await db.commit()
        self.send.side_effect = MaxApiError(code=403, raw={"code": "chat.denied"})
        await worker.dispatch()
        self.assertEqual(sorted(n.status for n in await self.notifications()), ["cancelled", "failed"],
                         "первое не ушло, остальные из очереди сняты")
        await worker.queue_reminders(MORNING + timedelta(days=1))
        self.assertEqual(len(await self.notifications()), 2, "неактивным не пишем")
        async with SessionLocal() as db:
            self.assertFalse((await db.execute(select(User))).scalar_one().is_active)
        client = await self.api()
        with patch("app.api.routes.users.validate_max_init_data", lambda data: {"id": MAX_USER}):
            self.assertEqual((await client.post("/api/user/auth", json={"initData": "x"})).status_code, 200)
        async with SessionLocal() as db:
            self.assertTrue((await db.execute(select(User))).scalar_one().is_active)

    async def test_one_company_one_user(self):
        from databases.businesses_db import CompanyTaken
        async with SessionLocal() as db:
            other = User(max_user_id=222)
            db.add(other)
            await db.flush()
            with self.assertRaises(CompanyTaken):
                await save_business(db, other.id, RECORD)
            await db.commit()
        client = await self.api(222)
        response = await client.post("/api/session", json={"inn": INN})
        self.assertEqual(response.status_code, 409)

    async def test_profile_form_saves_only_corrections(self):
        from bot import message_handler as bot_flow
        async with SessionLocal() as db:
            user = (await db.execute(select(User))).scalar_one()
            user.awaiting_mid, user.awaiting_since = "mid-card", datetime.now(MSK)
            await db.commit()
        client = await self.api()
        form = (await client.get("/api/company/profile")).json()
        self.assertEqual((form.get("regime"), form["okved"], form["region"], form["hasLicenses"], form["registryRegion"]),
                         (None, "41.20", "16", False, "Республика Татарстан"))
        self.assertNotIn("psn", [o["value"] for o in form["options"]["regimes"]], "патент — только у ИП")
        self.assertIn("cash_register", form["flags"])

        update = {"regime": "usn_income", "headcount": "16", "okved": "56.10", "region": "16", "hasLicenses": True,
                  "flags": {"cash_register": True, "marked_goods": None}, "kpp": "165001001",
                  "address": "423800, г. Набережные Челны, пр. Мира, д. 1", "directorPosition": "Генеральный директор",
                  "directorName": "ИВАНОВ ИВАН ИВАНОВИЧ"}
        for bad in ({"okved": "ресторан"}, {"regime": "psn"}, {"region": "00"}, {"headcount": "7"}, {"kpp": "12345"}):
            self.assertEqual((await client.put("/api/company/profile", json=update | bad)).status_code, 422, bad)
        next_step = AsyncMock(return_value=False)
        with patch.object(bot_flow, "next_step", next_step):
            response = await client.put("/api/company/profile", json=update)
            self.assertEqual(response.status_code, 200)
            self.assertIn("56.10", response.json()["okved"])
            next_step.assert_awaited_once_with(MAX_USER)
            self.deleted.assert_awaited_once_with("mid-card")
            client = await self.api()
            await client.put("/api/company/profile", json=update)
            next_step.assert_awaited_once()

        async with SessionLocal() as db:
            answers = await db.get(BusinessProfile, INN)
            profile = await worker.load_profile(db, INN)
        self.assertEqual((answers.okved_main, answers.region_code, answers.has_licenses), ("56.10", None, True),
                         "совпадающее с реестром не храним — обновления реестра подтянутся")
        self.assertEqual((profile.okved_main, profile.region_code, profile.headcount, profile.flags),
                         ("56.10", "16", 16, {"cash_register": True}))
        self.assertTrue(await self.events(type="deadline"), "обязанности пересчитаны")
        self.assertEqual((answers.kpp, answers.director_position), ("165001001", "Генеральный директор"))
        form = (await client.get("/api/company/profile")).json()
        self.assertEqual((form["kpp"], form["directorName"]), ("165001001", "ИВАНОВ ИВАН ИВАНОВИЧ"))
        self.assertEqual((await client.get("/api/company")).json()["kpp"], "165001001")

    async def test_registration_card_waits_for_answer(self):
        from bot import message_handler as bot_flow
        say = AsyncMock(return_value="mid-card")
        with patch.object(bot_flow, "say", say), patch.object(bot_flow, "bot_app", AsyncMock(return_value=BOT_APP)):
            async with SessionLocal() as db:
                business = await db.get(Business, INN)
            await bot_flow.show_found(MAX_USER, business)
        _, text, kb = say.await_args.args
        self.assertTrue(text.startswith("Нашли вашу компанию:"))
        self.assertIn("Режим налогообложения: в открытых реестрах его нет — спросим", text)
        self.assertTrue(text.endswith("Всё верно?"))
        self.assertEqual([(b.type, b.payload) for row in kb.payload.buttons for b in row],
                         [("callback", "reg:ok"), ("open_app", "profile_edit")])
        async with SessionLocal() as db:
            self.assertEqual((await db.execute(select(User))).scalar_one().awaiting_mid, "mid-card")

    async def test_bot_asks_profile_questions_then_shows_obligations(self):
        from bot import message_handler as bot_flow
        say = AsyncMock(return_value="mid-q")
        with patch.object(bot_flow, "say", say), patch.object(bot_flow, "bot_app", AsyncMock(return_value=BOT_APP)):
            await bot_flow.next_step(MAX_USER)
            self.assertIn("режим налогообложения", say.await_args.args[1])
            await self.answer("usn_income", None)
            await bot_flow.next_step(MAX_USER)
            self.assertIn("Сколько у вас сотрудников", say.await_args.args[1])
            await self.answer("usn_income", 1)
            await bot_flow.next_step(MAX_USER)
        summary = say.await_args.args[1]
        self.assertIn("Радар настроен: 10 обязанностей", summary)
        self.assertIn("Ближайший срок — ", summary)
        self.assertIn("Список, календарь и документы — в приложении", summary)
        async with SessionLocal() as db:
            card = bot_flow.company_card(await db.get(Business, INN), await worker.load_profile(db, INN), "реестр МСП, открытые данные ФНС на 10.09.2026")
        self.assertIn("ОКВЭД: 41.20 — Строительство жилых и нежилых зданий", card)
        self.assertIn("Микропредприятие · 1–15 человек", card)
        self.assertIn("Регион: Республика Татарстан", card)
        self.assertIn("Лицензии: нет", card)

    async def test_registry_headcount_skips_question(self):
        from bot import message_handler as bot_flow
        async with SessionLocal() as db:
            (await db.get(Business, INN)).employees_num = 12
            await db.commit()
        say = AsyncMock(return_value="mid-q")
        with patch.object(bot_flow, "say", say), patch.object(bot_flow, "bot_app", AsyncMock(return_value=BOT_APP)):
            await self.answer("usn_income", None)
            await bot_flow.next_step(MAX_USER)
        summary = say.await_args.args[1]
        self.assertIn("Радар настроен", summary, "о численности не спрашиваем — она есть в реестре")
        async with SessionLocal() as db:
            profile = await worker.load_profile(db, INN)
            self.assertIn("Микропредприятие · 12 человек (ФНС)", bot_flow.company_card(await db.get(Business, INN), profile, "реестр МСП, открытые данные ФНС на 10.09.2026"))
        benefits = [b.title for b in profile.headcount_benefits()]
        self.assertTrue(any(title.startswith("Можно применять УСН") for title in benefits))
        self.assertFalse(any("АУСН" in title for title in benefits), "12 человек — больше порога АУСН")
        self.assertEqual((profile.headcount, profile.headcount_exact, profile.has_employees), (12, True, True))

        await self.answer("usn_income", 36)
        async with SessionLocal() as db:
            profile = await worker.load_profile(db, INN)
        self.assertEqual((profile.headcount, profile.headcount_exact), (36, False))


if __name__ == "__main__":
    unittest.main()
