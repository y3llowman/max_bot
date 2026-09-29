from __future__ import annotations

import io
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date

from docx import Document
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from databases.businesses_db import Business
from notifications.models import BusinessProfile, RegistrySnapshot
from radar.deadlines import region_name
from radar.render import as_date, company_name, date_ru, fio

TITLES = {
    "notice": "Уведомление об исчисленных суммах",
    "quota_order": "Приказ о квотируемых рабочих местах",
    "enp_payment": "Реквизиты платёжки на ЕНП",
    "report_brief": "Памятка к отчёту",
}

HINTS = {
    "notice": "Проверьте реквизиты, впишите суммы и отправьте.",
    "quota_order": "Проверьте реквизиты, впишите число мест и подпишите.",
    "enp_payment": "Перенесите реквизиты в платёжку в банке и впишите сумму.",
    "report_brief": "Это не бланк: сам отчёт сформируйте в бухгалтерской программе или личном кабинете.",
}

KBK = {
    "ndfl": ("НДФЛ", "18210102010011000110"),
    "insurance": ("Страховые взносы", "18210201000011000160"),
    "usn_income": ("Аванс по УСН «доходы»", "18210501011011000110"),
    "usn_ie": ("Аванс по УСН «доходы минус расходы»", "18210501021011000110"),
    "ndfl_ip": ("Аванс по НДФЛ ИП — КБК зависит от суммы дохода, сверьте в личном кабинете ИП", None),
}

BLANK = "________"


async def requisites(db: AsyncSession, inn: str) -> dict:
    business = await db.get(Business, inn)
    profile = await db.get(BusinessProfile, inn)
    snapshot = (await db.execute(
        select(RegistrySnapshot.data)
        .where(RegistrySnapshot.inn == inn, RegistrySnapshot.source == "egrul")
        .order_by(RegistrySnapshot.fetched_at.desc()).limit(1)
    )).scalar_one_or_none() or {}
    director = snapshot.get("director") or {}
    director_name = " ".join(filter(None, (director.get("surname"), director.get("name"), director.get("patronymic"))))
    return {
        "name": company_name(snapshot.get("full_name") or business.name),
        "inn": inn,
        "kpp": snapshot.get("kpp"),
        "ogrn": business.ogrn,
        "address": snapshot.get("address"),
        "director_position": (director.get("position") or "Руководитель").capitalize(),
        "director_name": fio(director_name) or None,
        "region": region_name(business.region_code),
        "tax_regime": profile.tax_regime if profile else None,
    }


def _requisites_table(doc: Document, req: dict) -> None:
    rows = [("Организация", req["name"]), ("ИНН", req["inn"]), ("КПП", req["kpp"]),
            ("ОГРН", req["ogrn"]), ("Адрес", req["address"])]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value or BLANK


def _notice(doc: Document, payload: dict, due: date, req: dict) -> None:
    doc.add_heading("Уведомление об исчисленных суммах налогов (КНД 1110355)", level=1)
    doc.add_paragraph(f"{payload['title']} · {payload['period']}. Подать до {date_ru(due)}.")
    doc.add_paragraph("Черновик: реквизиты заполнены из реестров ФНС. Впишите ОКТМО и суммы, затем "
                      "отправьте уведомление через оператора ЭДО или личный кабинет налогоплательщика.")
    _requisites_table(doc, req)
    doc.add_paragraph()
    kinds = {"ndfl_notice": ["ndfl", "insurance"], "ip_ndfl_notice": ["ndfl_ip"]}.get(payload["code"], [req["tax_regime"]])
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Платёж", "КБК", "ОКТМО", "Сумма, ₽")):
        cell.text = text
    for kind in kinds:
        label, kbk = KBK.get(kind, ("Налог", BLANK))
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text, cells[3].text = label, kbk or BLANK, BLANK, BLANK


def _quota_order(doc: Document, payload: dict, due: date, req: dict) -> None:
    doc.add_paragraph(req["name"])
    doc.add_heading("Приказ № ____", level=1)
    doc.add_paragraph("«___» __________ 20__ г.")
    doc.add_paragraph("О выделении рабочих мест для трудоустройства инвалидов в счёт квоты").runs[0].bold = True
    region = f"закона субъекта РФ ({req['region']})" if req["region"] else "закона субъекта РФ"
    doc.add_paragraph(f"В соответствии со статьёй 38 Федерального закона от 12.12.2023 № 565-ФЗ «О занятости "
                      f"населения в Российской Федерации» и {region} о квоте для приёма на работу инвалидов")
    doc.add_paragraph("ПРИКАЗЫВАЮ:")
    for item in (
        f"Выделить в счёт квоты {BLANK} рабочих мест для трудоустройства инвалидов: "
        "размер квоты — процент от среднесписочной численности, установленный законом региона.",
        "Утвердить перечень квотируемых рабочих мест (приложение к приказу).",
        "Ежемесячно, не позднее 10-го числа, представлять сведения о выполнении квоты по форме № 7 "
        f"на портале «Работа России». Ответственный: {BLANK}.",
        "Контроль за исполнением приказа оставляю за собой.",
    ):
        doc.add_paragraph(item, style="List Number")
    doc.add_paragraph()
    doc.add_paragraph(f"{req['director_position']} ____________ {req['director_name'] or BLANK}")


ENP_CHECKED = date(2026, 9, 28)
ENP_RECIPIENT = (
    ("13", "Банк получателя", "ОКЦ № 7 ГУ Банка России по ЦФО//УФК по Тульской области, г Тула"),
    ("14", "БИК банка получателя", "017003983"),
    ("15", "Счёт банка получателя", "40102810445370000059"),
    ("16", "Получатель", "Казначейство России (ФНС России)"),
    ("17", "Казначейский счёт", "03100643000000018500"),
    ("61", "ИНН получателя", "7727406020"),
    ("103", "КПП получателя", "770701001"),
    ("104", "КБК", "18201061201010000510"),
    ("105", "ОКТМО", "0"),
)


def _fields_table(doc: Document, rows: list[tuple[str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Поле", "Реквизит", "Значение")):
        cell.text = text
    for number, label, value in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = number, label, value or BLANK


def _enp_payment(doc: Document, payload: dict, due: date, req: dict) -> None:
    doc.add_heading("Платёжное поручение на единый налоговый платёж", level=1)
    doc.add_paragraph(f"{payload['title']} · {payload['period']}. Заплатить до {date_ru(due)}.")
    doc.add_paragraph("Деньги поступят на единый налоговый счёт, а ФНС распределит их по налогам "
                      "из уведомлений и деклараций. Поэтому КБК один для всех налогов, а основание "
                      "платежа и налоговый период не указываются.")
    payer = [
        ("101", "Статус плательщика", "01"),
        ("60", "ИНН плательщика", req["inn"]),
        ("102", "КПП плательщика", "0"),
        ("8", "Плательщик", req["name"]),
        ("7", "Сумма", BLANK),
        ("24", "Назначение платежа", "Единый налоговый платёж"),
    ]
    _fields_table(doc, payer + list(ENP_RECIPIENT))
    doc.add_paragraph()
    doc.add_paragraph("Сумма — из уведомлений и деклараций этого месяца. Если платите за ИП "
                      "«за себя», сумма фиксированных взносов на год — на сайте ФНС.")
    if req["kpp"]:
        doc.add_paragraph(f"Если банк не принимает «0» в поле 102, укажите КПП организации ({req['kpp']}): "
                          "приказ Минфина № 107н допускает оба значения.")
    doc.add_paragraph(f"Реквизиты сверены {date_ru(ENP_CHECKED)}. Перед оплатой сверьтесь с "
                      "сайтом ФНС: при смене реквизитов платёж может попасть в невыясненные.")


PeriodFn = Callable[[date], tuple[str | None, int]]


def _annual(n: date) -> tuple[str | None, int]:
    return "34", n.year - 1


def _cumulative(annual_month: int) -> PeriodFn:
    def period(n: date) -> tuple[str | None, int]:
        if n.month == annual_month:
            return "34", n.year - 1
        return {4: "21", 7: "31", 10: "33"}.get(n.month), n.year
    return period


def _vat(n: date) -> tuple[str | None, int]:
    if n.month == 1:
        return "24", n.year - 1
    return {4: "21", 7: "22", 10: "23"}.get(n.month), n.year


def _year_only(n: date) -> tuple[str | None, int]:
    return None, n.year - 1


@dataclass(frozen=True)
class Report:
    form: str
    knd: str | None
    period: PeriodFn | None
    needs: tuple[str, ...]


FNS_TOOL_LE = ("бухгалтерская программа, личный кабинет налогоплательщика или бесплатная "
               "программа ФНС «Налогоплательщик ЮЛ»")
FNS_TOOL_IP = "личный кабинет ИП на сайте ФНС или бухгалтерская программа"
SFR_TOOL = "кабинет страхователя на сайте Социального фонда или бухгалтерская программа"

REPORTS: dict[str, Report] = {
    "usn_decl": Report(
        "Декларация по УСН", "1152017", _annual,
        ("Книга учёта доходов и расходов (КУДиР) за год.",
         "Суммы авансов за I квартал, полугодие и 9 месяцев.",
         "Уплаченные страховые взносы — на УСН «доходы» они уменьшают налог.")),
    "psfl": Report(
        "Персонифицированные сведения о физических лицах", "1151162", None,
        ("Список физлиц, получивших выплаты за месяц: ФИО, СНИЛС, ИНН.",
         "Сумма выплат каждому, включая исполнителей по договорам ГПХ.")),
    "rsv": Report(
        "Расчёт по страховым взносам", "1151111", _cumulative(1),
        ("Выплаты каждому работнику по месяцам периода.",
         "Начисленные взносы и применённый тариф.",
         "СНИЛС, ИНН и категории застрахованных лиц.")),
    "ndfl6": Report(
        "Расчёт 6-НДФЛ", "1151100", _cumulative(2),
        ("Доходы работников и исполнителей нарастающим итогом с начала года.",
         "Исчисленный, удержанный и перечисленный НДФЛ.",
         "Даты выплат и удержания налога.")),
    "efs1": Report(
        "ЕФС-1 (Социальный фонд России)", None, None,
        ("Начисленные и уплаченные взносы на травматизм за период.",
         "В январе — периоды работы и стаж каждого сотрудника за прошлый год.")),
    "efs1_ausn": Report(
        "ЕФС-1, подраздел 1.2 (стаж)", None, _year_only,
        ("Периоды работы каждого сотрудника за прошлый год.",
         "Условия работы: коды особых условий труда, если есть.")),
    "buh": Report(
        "Бухгалтерская (финансовая) отчётность", None, _year_only,
        ("Остатки по счетам на 31 декабря.",
         "Доходы и расходы за год.",
         "Для малого бизнеса, как правило, достаточно упрощённых форм баланса и отчёта "
         "о финансовых результатах.")),
    "vat_decl": Report(
        "Декларация по НДС", "1151001", _vat,
        ("Книга продаж и книга покупок за квартал.",
         "Счета-фактуры, по которым заявляете вычеты.",
         "Восстановленный НДС, если был.")),
    "profit_decl": Report(
        "Декларация по налогу на прибыль", "1151006", _cumulative(3),
        ("Доходы и расходы налогового учёта нарастающим итогом.",
         "Авансовые платежи, начисленные за период.")),
    "ip_3ndfl": Report(
        "Декларация 3-НДФЛ", "1151020", _annual,
        ("Доходы от предпринимательской деятельности за год.",
         "Документы на профессиональные вычеты — или норматив 20% без документов.",
         "Уплаченные авансовые платежи по НДФЛ.")),
}


def report_period(code: str, payload: dict, due: date) -> tuple[str | None, int | None]:
    report = REPORTS[code]
    if report.period is None:
        return None, None
    nominal = as_date(payload.get("original")) or due
    return report.period(nominal)


def _report_brief(doc: Document, payload: dict, due: date, req: dict) -> None:
    report = REPORTS[payload["code"]]
    is_ip = len(req["inn"]) == 12
    period_code, year = report_period(payload["code"], payload, due)
    doc.add_heading(f"{report.form} — памятка к отчёту", level=1)
    doc.add_paragraph(f"{payload['period']}. Сдать до {date_ru(due)}.")

    rows = [("Организация" if not is_ip else "Индивидуальный предприниматель", req["name"]),
            ("ИНН", req["inn"])]
    if not is_ip:
        rows.append(("КПП", req["kpp"]))
    rows += [("ОГРНИП" if is_ip else "ОГРН", req["ogrn"])]
    if report.knd:
        rows.append(("Форма по КНД", report.knd))
    if year:
        rows.append(("Отчётный год", str(year)))
    if period_code:
        rows.append(("Код периода", period_code))
    rows += [("Куда", payload["where"]), ("Формат", payload["format"])]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value or BLANK

    doc.add_heading("Что подготовить", level=2)
    for item in report.needs:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Порядок", level=2)
    for step in payload["how"]:
        if "Подготовить документ" not in step:
            doc.add_paragraph(step, style="List Number")
    doc.add_heading("Если не успеть", level=2)
    doc.add_paragraph(payload["penalty"])
    doc.add_paragraph(f"Основание: {payload['basis']}.")

    tool = SFR_TOOL if payload["code"] == "efs1" else FNS_TOOL_IP if is_ip else FNS_TOOL_LE
    doc.add_paragraph(f"Это памятка, а не бланк. Отчёт сформируйте по утверждённому формату: {tool}.")


BUILDERS = {
    "notice": _notice,
    "quota_order": _quota_order,
    "enp_payment": _enp_payment,
    "report_brief": _report_brief,
}


CAPTIONS = {
    "notice": "📄 {doc} — черновик к задаче «{title}» ({period}). ",
    "quota_order": "📄 {doc} — черновик к задаче «{title}» ({period}). ",
    "enp_payment": "📄 {doc} для задачи «{title}» ({period}). ",
    "report_brief": "📄 {doc}: «{title}» ({period}). ",
}


def caption(payload: dict) -> str:
    code = payload["document"]
    head = CAPTIONS[code].format(doc=TITLES[code], title=payload["title"], period=payload["period"])
    return head + HINTS[code]


def build(payload: dict, due: date, req: dict) -> tuple[str, bytes]:
    code = payload["document"]
    doc = Document()
    BUILDERS[code](doc, payload, due, req)
    buffer = io.BytesIO()
    doc.save(buffer)
    prefix = code if code in ("notice", "quota_order") else payload["code"]
    return f"{prefix}_{req['inn']}_{due:%Y%m%d}.docx", buffer.getvalue()
